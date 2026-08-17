# tool/doc_tool.py
from docx.oxml.ns import qn
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
import os
from datetime import datetime
dir = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
dir_path = os.path.join(dir, "outputs")
# 生成 Word 文档
def generate_word(content: str, title: str = "医疗报告", output_dir: str = dir_path) -> str:
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, f"{title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")

    doc = Document()

    # 设置标题
    title_run = doc.add_heading(title, level=1).runs[0]
    title_run.font.name = "SimSun"
    title_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

    # 设置正文
    para = doc.add_paragraph(content)
    run = para.runs[0]
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    run.font.size = Pt(12)

    doc.save(file_path)
    return file_path


# 生成 PDF 文档
def generate_pdf(content: str, title: str = "医疗报告", output_dir: str = dir_path) -> str:
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, f"{title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")

    c = canvas.Canvas(file_path, pagesize=A4)
    width, height = A4

    # 标题
    c.setFont("Helvetica-Bold", 16)
    c.drawString(100, height - 100, title)

    # 内容
    c.setFont("Helvetica", 12)
    y = height - 150
    for line in content.split("\n"):
        c.drawString(80, y, line)
        y -= 20
        if y < 50:  # 自动分页
            c.showPage()
            c.setFont("Helvetica", 12)
            y = height - 100

    c.save()
    return file_path

if __name__ == '__main__':
    print(generate_word("hello world"))